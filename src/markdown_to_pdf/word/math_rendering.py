"""
LaTeX → OMML (Word-native math) conversion.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree as _lxml_etree
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from markdown_to_pdf.constants import (
    rgb_docx,
    COLOR_EM,
    FONT_BODY_PT,
    WordFonts,
    XSLT_PATH,
)
from markdown_to_pdf.word.ooxml import (
    no_space_before,
    set_para_space_after,
    set_run_font,
)

_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MML_TRANSFORM = None


def _get_mml_transform():
    """Lazily load the Microsoft MML2OMML XSLT transform."""
    global _MML_TRANSFORM
    xslt_path = Path(XSLT_PATH)
    if _MML_TRANSFORM is None and xslt_path.exists():
        try:
            xslt_doc = _lxml_etree.parse(str(xslt_path))
            _MML_TRANSFORM = _lxml_etree.XSLT(xslt_doc)
        except Exception:
            pass
    return _MML_TRANSFORM


def latex_to_omml(latex_src: str):
    """Convert LaTeX to a Word OMML ``<m:oMath>`` element.  Returns ``None`` on failure."""
    try:
        import latex2mathml.converter
        mml_str  = latex2mathml.converter.convert(latex_src)
        transform = _get_mml_transform()
        if transform is None:
            return None
        mml_doc = _lxml_etree.fromstring(mml_str.encode())
        result  = transform(mml_doc)
        root    = result.getroot()
        if root is None or len(root) == 0:
            return None
        return root
    except Exception:
        return None


def insert_math_block(doc, latex_src: str) -> None:
    """Insert a centred display equation into the document."""
    para = doc.add_paragraph()
    no_space_before(para)
    set_para_space_after(para, 10)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    omml = latex_to_omml(latex_src)
    if omml is not None:
        oMathPara = _lxml_etree.Element(f"{{{_MATH_NS}}}oMathPara")
        oMathPara.append(omml)
        para._p.append(oMathPara)
    else:
        run = para.add_run(f"$$ {latex_src} $$")
        run.italic = True
        set_run_font(run, WordFonts.BODY)
        run.font.size = Pt(FONT_BODY_PT)
        run.font.color.rgb = rgb_docx(COLOR_EM)


def insert_math_inline(para, latex_src: str) -> None:
    """Append an inline equation run to an existing paragraph."""
    omml = latex_to_omml(latex_src)
    if omml is not None:
        para._p.append(omml)
    else:
        run = para.add_run(f"${latex_src}$")
        run.italic = True
        set_run_font(run, WordFonts.BODY)
        run.font.size = Pt(FONT_BODY_PT)
        run.font.color.rgb = rgb_docx(COLOR_EM)
