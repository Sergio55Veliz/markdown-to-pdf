"""
Word rendering — inline runs, block elements, blockquotes, and emoji.

This module contains all functions that convert BeautifulSoup nodes into
Word ``Paragraph`` / ``Run`` / ``Table`` elements.
"""

from __future__ import annotations

import copy
import re
import urllib.parse

from lxml import etree as _lxml_etree
from bs4 import NavigableString, Tag
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from markdown_to_pdf.constants import (
    rgb_docx,
    rgb_hex_upper,
    WordFonts,
    FONT_BODY_PT,
    FONT_CODE_PT,
    PARA_SPACING_PT,
    HEADING_BOLD,
    HEADING_COLORS,
    EMOJI_RE,
    BLACK_CIRCLE,
    CIRCLE_EMOJI,
    EMOJI_COLOR,
    COLOR_BODY,
    COLOR_BQ_BG,
    COLOR_BQ_BORDER,
    COLOR_BQ_TEXT,
    COLOR_CODE_TEXT,
    COLOR_HR,
    COLOR_LINK,
    COLOR_PRE_BG,
    COLOR_PRE_TEXT,
    COLOR_TH_BG,
    COLOR_TH_TEXT,
    COLOR_TD_BORDER,
    COLOR_TD_ALT_BG,
)
from markdown_to_pdf.word.ooxml import (
    set_para_shading,
    set_para_border_left,
    set_para_indent,
    no_space_before,
    set_para_space_after,
    set_cell_shading,
    set_cell_border,
    set_run_font,
    add_bookmark,
)
from markdown_to_pdf.word.math_rendering import insert_math_block, insert_math_inline
from markdown_to_pdf.word.document import setup_document


# ── Heading slug state (reset per document via reset_rendering_state) ─────────

_slug_counts: dict[str, int] = {}
_pending_anchor: str | None = None


def reset_rendering_state() -> None:
    """Reset per-document rendering state.  Call before each conversion."""
    global _slug_counts, _pending_anchor
    _slug_counts = {}
    _pending_anchor = None


def _heading_slug(text: str) -> str:
    """Compute a GitHub-style anchor slug and register it globally."""
    global _slug_counts
    base = re.sub(r'\s+', '-', re.sub(r'[^\w\s-]', '', text.lower().strip()))
    count = _slug_counts.get(base, 0)
    _slug_counts[base] = count + 1
    return base if count == 0 else f"{base}-{count}"


# ── Emoji helpers ─────────────────────────────────────────────────────────────

def _split_by_emoji(text: str) -> list[tuple[str, str]]:
    """Split text into ``('text', …)`` and ``('emoji', …)`` segments."""
    result: list[tuple[str, str]] = []
    last = 0
    for m in EMOJI_RE.finditer(text):
        if m.start() > last:
            result.append(("text", text[last:m.start()]))
        result.append(("emoji", m.group(1)))
        last = m.end()
    if last < len(text):
        result.append(("text", text[last:]))
    return result


# ── Naked anchor detection ────────────────────────────────────────────────────

def is_naked_anchor(node) -> str | None:
    """Return the ``id`` if node is a bare ``<a id="...">`` anchor, else ``None``."""
    if node.name == "a" and node.get("id") and not node.get("href") \
            and not node.get_text(strip=True):
        return node.get("id")
    if node.name == "p":
        meaningful = [
            c for c in node.children
            if not (isinstance(c, NavigableString) and not c.strip())
        ]
        if len(meaningful) == 1:
            child = meaningful[0]
            if hasattr(child, "name") and child.name == "a" \
                    and child.get("id") and not child.get("href") \
                    and not child.get_text(strip=True):
                return child.get("id")
    return None


# ── Run-level formatting ─────────────────────────────────────────────────────

def _apply_run_style(run, *, bold=False, italic=False, strike=False,
                     color=None, font_name=None, font_size=None,
                     highlight_color=None) -> None:
    if font_name:
        set_run_font(run, font_name)
    if font_size is not None:
        run.font.size = Pt(font_size)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if strike:
        run.font.strike = True
    if color:
        run.font.color.rgb = rgb_docx(color)
    if highlight_color:
        run.font.color.rgb = rgb_docx(COLOR_CODE_TEXT)


# ── Inline HTML → runs ───────────────────────────────────────────────────────

def inline_nodes_to_runs(para, node, bold=False, italic=False, strike=False,
                         color=None, is_code_block=False) -> None:
    """Recursively walk an inline BeautifulSoup node and add formatted runs."""
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        fn         = WordFonts.MONO if is_code_block else WordFonts.BODY
        fpt        = FONT_CODE_PT if is_code_block else FONT_BODY_PT
        base_color = color or COLOR_BODY
        for kind, val in _split_by_emoji(text):
            if kind == "text":
                run = para.add_run(val)
                _apply_run_style(
                    run, bold=bold, italic=italic, strike=strike,
                    color=base_color, font_name=fn, font_size=fpt,
                )
            elif val in CIRCLE_EMOJI:
                run = para.add_run(BLACK_CIRCLE)
                _apply_run_style(
                    run, bold=bold, font_name=fn, font_size=fpt,
                    color=CIRCLE_EMOJI[val],
                )
            else:
                emoji_color = EMOJI_COLOR.get(val, base_color)
                run = para.add_run(val)
                _apply_run_style(
                    run, bold=bold, font_name=WordFonts.EMOJI,
                    font_size=fpt, color=emoji_color,
                )
        return

    tag = node.name
    new_bold   = bold   or tag in ("strong", "b")
    new_italic = italic or tag in ("em", "i")
    new_strike = strike or tag == "del"
    new_color  = color

    if tag == "code":
        run = para.add_run(node.get_text())
        _apply_run_style(
            run, bold=bold, italic=italic,
            color=COLOR_CODE_TEXT, font_name=WordFonts.MONO,
            font_size=FONT_CODE_PT,
        )
        return

    if tag == "a":
        href = node.get("href", "")
        if href.startswith("#"):
            anchor = urllib.parse.unquote(href[1:]) # decodifica la url
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), anchor)
            p_elem = para._p
            idx_before = len(p_elem)
            for child in node.children:
                inline_nodes_to_runs(
                    para, child, bold=bold, italic=italic, strike=strike,
                    color=COLOR_LINK, is_code_block=is_code_block,
                )
            new_els = list(p_elem)[idx_before:]
            for el in new_els:
                p_elem.remove(el)
                hyperlink.append(el)
            p_elem.append(hyperlink)
            return
        new_color = COLOR_LINK

    if tag in ("strong", "b"):
        new_color = new_color

    if tag == "span" and "math-inline" in node.get("class", []):
        insert_math_inline(para, node.get_text())
        return

    for child in node.children:
        inline_nodes_to_runs(
            para, child,
            bold=new_bold, italic=new_italic, strike=new_strike,
            color=new_color or color, is_code_block=is_code_block,
        )


# ── Block-level rendering ────────────────────────────────────────────────────

def render_heading(doc: Document, tag: Tag, level: int) -> None:
    """Render a heading as a Word built-in Heading style."""
    global _pending_anchor
    hcolor = HEADING_COLORS[level]
    para = doc.add_paragraph(style=f"Heading {level}")
    for child in tag.children:
        inline_nodes_to_runs(para, child, bold=HEADING_BOLD[level], color=hcolor)

    for run in para.runs:
        fn = run.font.name
        set_run_font(run, WordFonts.MONO if fn == WordFonts.MONO else WordFonts.HEADING)

    slug = _heading_slug(tag.get_text(" ", strip=True))
    add_bookmark(para, slug)
    if _pending_anchor:
        add_bookmark(para, _pending_anchor)
        _pending_anchor = None


def render_paragraph(doc: Document, tag: Tag,
                     indent_twips: int = 0,
                     extra_color=None,
                     is_blockquote: bool = False) -> None:
    """Render a ``<p>`` tag as a styled body paragraph."""
    para = doc.add_paragraph()
    no_space_before(para)
    set_para_space_after(para, PARA_SPACING_PT)

    for child in tag.children:
        inline_nodes_to_runs(para, child, color=extra_color or COLOR_BODY)

    for run in para.runs:
        if not run.font.name:
            set_run_font(run, WordFonts.BODY)
        if not run.font.size:
            run.font.size = Pt(FONT_BODY_PT)

    if indent_twips:
        set_para_indent(para, indent_twips)
    if is_blockquote:
        set_para_shading(para, COLOR_BQ_BG)
        set_para_border_left(para, COLOR_BQ_BORDER)


def render_code_block(doc: Document, tag: Tag) -> None:
    """Render a ``<pre><code>`` fenced code block."""
    code_tag = tag.find("code")
    text = code_tag.get_text() if code_tag else tag.get_text()
    text = text.rstrip("\n")

    para = doc.add_paragraph()
    no_space_before(para)
    set_para_space_after(para, 8)
    set_para_shading(para, COLOR_PRE_BG)
    set_para_indent(para, 180)

    run = para.add_run(text)
    set_run_font(run, WordFonts.MONO)
    run.font.size  = Pt(FONT_CODE_PT)
    run.font.color.rgb = rgb_docx(COLOR_PRE_TEXT)

    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"),     "240")
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)


def render_table(doc: Document, tag: Tag) -> None:
    """Render a Markdown table as a python-docx Table."""
    rows_tags = tag.find_all("tr")
    if not rows_tags:
        return
    num_cols = max(len(r.find_all(["th", "td"])) for r in rows_tags)
    if num_cols == 0:
        return

    tbl = doc.add_table(rows=len(rows_tags), cols=num_cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    for r_idx, row_tag in enumerate(rows_tags):
        cells_tags = row_tag.find_all(["th", "td"])
        is_header  = any(c.name == "th" for c in cells_tags)

        for c_idx, cell_tag in enumerate(cells_tags):
            if c_idx >= num_cols:
                break
            cell = tbl.cell(r_idx, c_idx)

            if is_header:
                set_cell_shading(cell, COLOR_TH_BG)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, COLOR_TD_ALT_BG)

            para = cell.paragraphs[0]
            para.clear()
            no_space_before(para)
            set_para_space_after(para, 4)

            txt_color = COLOR_TH_TEXT if is_header else COLOR_BODY
            for child in cell_tag.children:
                inline_nodes_to_runs(para, child, color=txt_color, bold=is_header)

            for run in para.runs:
                set_run_font(run, WordFonts.BODY)
                run.font.size = Pt(FONT_CODE_PT)
                if is_header:
                    run.font.color.rgb = rgb_docx(COLOR_TH_TEXT)
                    run.bold = True

            if not is_header:
                set_cell_border(cell, COLOR_TD_BORDER)


def render_list(doc: Document, tag: Tag, level: int = 0) -> None:
    """Render a ``<ul>`` or ``<ol>`` element, including nested lists."""
    is_ordered = tag.name == "ol"
    counter    = 1
    _BLOCK_IN_LI = {"p", "div", "table", "hr", "pre", "blockquote", "ul", "ol"}

    for child in tag.children:
        if isinstance(child, NavigableString):
            continue
        if child.name != "li":
            continue

        indent_left = 360 + level * 360
        indent_hang = 360
        indent_text = indent_left + indent_hang

        if is_ordered:
            prefix = f"{counter}.\t"
            counter += 1
        else:
            prefix = "•\t"

        para = doc.add_paragraph()
        no_space_before(para)
        set_para_space_after(para, 3)
        pPr = para._p.get_or_add_pPr()
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"),    str(indent_text))
        ind.set(qn("w:hanging"), str(indent_hang))
        pPr.append(ind)

        run_prefix = para.add_run(prefix)
        set_run_font(run_prefix, WordFonts.BODY)
        run_prefix.font.size      = Pt(FONT_BODY_PT)
        run_prefix.font.color.rgb = rgb_docx(COLOR_BODY)

        first_block_done = False

        for li_child in child.children:
            if isinstance(li_child, NavigableString):
                text = str(li_child)
                if text.strip():
                    run = para.add_run(text)
                    set_run_font(run, WordFonts.BODY)
                    run.font.size      = Pt(FONT_BODY_PT)
                    run.font.color.rgb = rgb_docx(COLOR_BODY)
                    first_block_done = True
                continue

            if li_child.name not in _BLOCK_IN_LI:
                inline_nodes_to_runs(para, li_child, color=COLOR_BODY)
                continue

            if li_child.name == "p":
                if not first_block_done:
                    for sub in li_child.children:
                        inline_nodes_to_runs(para, sub, color=COLOR_BODY)
                    first_block_done = True
                else:
                    cont = doc.add_paragraph()
                    no_space_before(cont)
                    set_para_space_after(cont, 3)
                    set_para_indent(cont, indent_text)
                    for sub in li_child.children:
                        inline_nodes_to_runs(cont, sub, color=COLOR_BODY)
                    for run in cont.runs:
                        if not run.font.size:
                            run.font.size = Pt(FONT_BODY_PT)

            elif li_child.name == "div" and "math-block" in li_child.get("class", []):
                insert_math_block(doc, li_child.get_text(strip=True))

            elif li_child.name in ("ul", "ol"):
                render_list(doc, li_child, level=level + 1)

            elif li_child.name == "table":
                render_table(doc, li_child)

            elif li_child.name == "hr":
                render_hr(doc)

            elif li_child.name == "pre":
                render_code_block(doc, li_child)

            elif li_child.name == "blockquote":
                render_blockquote(doc, li_child)

            else:
                text = li_child.get_text(" ", strip=True)
                if text:
                    cont = doc.add_paragraph(text)
                    set_para_space_after(cont, 3)

            first_block_done = True

        for run in para.runs:
            if not run.font.size:
                run.font.size = Pt(FONT_BODY_PT)


def render_hr(doc: Document) -> None:
    """Render a Markdown ``<hr>`` as a paragraph bottom border."""
    para = doc.add_paragraph()
    no_space_before(para)
    set_para_space_after(para, 12)
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), rgb_hex_upper(COLOR_HR))
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── Blockquote ────────────────────────────────────────────────────────────────

def _bq_estimate_height_emu(tag: Tag) -> int:
    """Estimate the rendered height of a blockquote's content in EMU."""
    FONT_PT      = FONT_BODY_PT - 0.5
    FONT_LIST_PT = FONT_BODY_PT
    LINE_H_PT    = FONT_PT      * 1.20
    LINE_H_LIST  = FONT_LIST_PT * 1.20
    SPACE_AFT_PT = 4.0
    SPACE_AFT_LI = 3.0
    MATH_LINES   = 2.5
    TABLE_ROW_PT = LINE_H_PT + SPACE_AFT_PT
    PAD_TOP_PT   = 45720 / 12700
    PAD_BOT_PT   = 45720 / 12700
    USABLE_W_CM    = 15.5 - 2 * (91440 / 360000)
    CHARS_PER_LINE = USABLE_W_CM * 28.35 / (FONT_PT      * 0.60)
    CHARS_PER_LIST = USABLE_W_CM * 28.35 / (FONT_LIST_PT * 0.60)

    def _count_lines(text: str, chars_per_line: float = CHARS_PER_LINE) -> float:
        chars = len(text.strip())
        return 0.0 if chars == 0 else max(1.0, chars / chars_per_line)

    def _li_primary_text(li_tag) -> str:
        parts = []
        for c in li_tag.children:
            if isinstance(c, NavigableString):
                parts.append(str(c))
            elif hasattr(c, "name") and c.name not in ("ul", "ol"):
                parts.append(c.get_text())
        return " ".join(parts)

    def _walk(node: Tag) -> float:
        total = 0.0
        for child in node.children:
            if isinstance(child, NavigableString):
                txt = str(child).strip()
                if txt:
                    total += _count_lines(txt) * LINE_H_PT + SPACE_AFT_PT
                continue
            if not hasattr(child, "name") or child.name is None:
                continue
            name = child.name
            if name == "p":
                total += _count_lines(child.get_text()) * LINE_H_PT + SPACE_AFT_PT
            elif name in ("ul", "ol"):
                for li in child.find_all("li", recursive=False):
                    li_lines = _count_lines(_li_primary_text(li), CHARS_PER_LIST)
                    total += li_lines * LINE_H_LIST + SPACE_AFT_LI
                    for sub in li.find_all(["ul", "ol"], recursive=False):
                        for sub_li in sub.find_all("li", recursive=False):
                            sub_lines = _count_lines(
                                _li_primary_text(sub_li), CHARS_PER_LIST
                            )
                            total += sub_lines * LINE_H_LIST + SPACE_AFT_LI
            elif name == "div" and "math-block" in child.get("class", []):
                total += MATH_LINES * LINE_H_PT + 10.0
            elif name == "table":
                rows = child.find_all("tr")
                total += len(rows) * TABLE_ROW_PT
            elif name == "hr":
                total += 12.0
            elif name == "pre":
                lines = child.get_text().count("\n") + 1
                total += lines * LINE_H_PT + 8.0
            elif name == "blockquote":
                total += _bq_estimate_height_emu(child) / 12700
        return total

    content_pt = _walk(tag)
    safety_pt = max(4.0, content_pt * 0.10)
    total_pt  = PAD_TOP_PT + content_pt + PAD_BOT_PT + safety_pt
    return int(total_pt * 12700)


# Blockquote counter, shared across calls within a document
_bq_counter: int = 100


def reset_bq_counter() -> None:
    global _bq_counter
    _bq_counter = 100


def render_blockquote(doc: Document, tag: Tag) -> None:
    """Render a ``<blockquote>`` as a rounded-rectangle inline text box."""
    global _bq_counter

    _NS = {
        "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
        "wp":  "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
        "w":   "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    }

    def _sub(parent, ns_key, local, **attribs):
        child_el = _lxml_etree.SubElement(parent, f"{{{_NS[ns_key]}}}{local}")
        for k, v in attribs.items():
            child_el.set(k, str(v))
        return child_el

    def _hex(c):
        return rgb_hex_upper(c)

    # 1. Render children into scratch document
    scratch = setup_document()
    for p in scratch.paragraphs:
        p._element.getparent().remove(p._element)

    for child in tag.children:
        if isinstance(child, NavigableString):
            text = str(child).strip()
            if text:
                para = scratch.add_paragraph()
                no_space_before(para)
                set_para_space_after(para, 4)
                run = para.add_run(text)
                set_run_font(run, WordFonts.BODY)
                run.font.size = Pt(FONT_BODY_PT - 0.5)
                run.font.color.rgb = rgb_docx(COLOR_BQ_TEXT)
        elif child.name == "p":
            para = scratch.add_paragraph()
            no_space_before(para)
            set_para_space_after(para, 4)
            for sub in child.children:
                inline_nodes_to_runs(para, sub, color=COLOR_BQ_TEXT)
            for run in para.runs:
                set_run_font(run, WordFonts.BODY)
                run.font.size = Pt(FONT_BODY_PT - 0.5)
        elif child.name in ("ul", "ol"):
            render_list(scratch, child)
        elif child.name == "blockquote":
            render_blockquote(scratch, child)
        elif child.name == "div" and "math-block" in child.get("class", []):
            insert_math_block(scratch, child.get_text(strip=True))
        elif child.name == "table":
            render_table(scratch, child)
        elif child.name == "hr":
            render_hr(scratch)
        elif child.name == "pre":
            render_code_block(scratch, child)

    scratch_body = scratch.element.body

    # 2. Dimensions
    WIDTH_EMU  = int(15.5 * 360000)
    HEIGHT_EMU = _bq_estimate_height_emu(tag)

    RADIUS_EMU = int(6 * 12700)
    adj_val    = min(50000, int(RADIUS_EMU * 200000 // HEIGHT_EMU))

    # 3. Build DrawingML shape
    graphic     = _lxml_etree.Element(f"{{{_NS['a']}}}graphic")
    graphicData = _sub(
        graphic, "a", "graphicData",
        uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    )
    wsp = _sub(graphicData, "wps", "wsp")

    cNvSpPr = _sub(wsp, "wps", "cNvSpPr", txBx="1")
    _sub(cNvSpPr, "a", "spLocks", noChangeArrowheads="1")

    spPr = _sub(wsp, "wps", "spPr")
    xfrm = _sub(spPr, "a", "xfrm")
    _sub(xfrm, "a", "off", x="0", y="0")
    _sub(xfrm, "a", "ext", cx=str(WIDTH_EMU), cy=str(HEIGHT_EMU))

    prstGeom = _sub(spPr, "a", "prstGeom", prst="roundRect")
    avLst    = _sub(prstGeom, "a", "avLst")
    _sub(avLst, "a", "gd", name="adj", fmla=f"val {adj_val}")

    solidFill = _sub(spPr, "a", "solidFill")
    _sub(solidFill, "a", "srgbClr", val=_hex(COLOR_BQ_BG))

    ln     = _sub(spPr, "a", "ln", w="12700", cap="flat", cmpd="sng", algn="ctr")
    lnFill = _sub(ln, "a", "solidFill")
    _sub(lnFill, "a", "srgbClr", val=_hex(COLOR_BQ_BORDER))
    _sub(ln, "a", "prstDash", val="solid")

    # 4. Text box content
    txbx        = _sub(wsp, "wps", "txbx")
    txbxContent = _sub(txbx, "w", "txbxContent")

    for child_el in list(scratch_body):
        local = _lxml_etree.QName(child_el.tag).localname
        if local in ("p", "tbl"):
            txbxContent.append(copy.deepcopy(child_el))

    bodyPr = _sub(wsp, "wps", "bodyPr",
                  rot="0", spcFirstLastPara="0",
                  vertOverflow="overflow", horzOverflow="overflow",
                  vert="horz", wrap="square",
                  lIns="91440", tIns="45720", rIns="91440", bIns="45720",
                  anchor="t", anchorCtr="0")
    _sub(bodyPr, "a", "spAutoFit")

    # 5. Inline drawing wrapper
    _bq_counter += 1
    bq_id = _bq_counter

    drawing = _lxml_etree.Element(f"{{{_NS['w']}}}drawing")
    inline  = _lxml_etree.SubElement(
        drawing, f"{{{_NS['wp']}}}inline",
        distT="0", distB="0", distL="0", distR="0",
    )
    _sub(inline, "wp", "extent", cx=str(WIDTH_EMU), cy=str(HEIGHT_EMU))
    _sub(inline, "wp", "effectExtent", l="0", t="0", r="0", b="0")
    _sub(inline, "wp", "docPr", id=str(bq_id), name=f"Blockquote{bq_id}")
    cNvGFPr = _sub(inline, "wp", "cNvGraphicFramePr")
    _sub(cNvGFPr, "a", "graphicFrameLocks", noChangeAspect="1")
    inline.append(graphic)

    # 6. Insert into document
    host_para = doc.add_paragraph()
    no_space_before(host_para)
    set_para_space_after(host_para, 12)
    run_el = OxmlElement("w:r")
    run_el.append(drawing)
    host_para._p.append(run_el)


# ── Top-level block dispatcher ────────────────────────────────────────────────

def walk_block(doc: Document, node) -> None:
    """Dispatch a top-level BeautifulSoup block element to the correct renderer."""
    global _pending_anchor

    if isinstance(node, NavigableString):
        text = str(node).strip()
        if text:
            para = doc.add_paragraph(text)
            set_para_space_after(para, PARA_SPACING_PT)
        return

    tag_name = node.name
    if tag_name is None:
        return

    anchor_id = is_naked_anchor(node)
    if anchor_id is not None:
        _pending_anchor = anchor_id
        return

    if tag_name in ("h1", "h2", "h3", "h4", "h5", "h6"):
        render_heading(doc, node, int(tag_name[1]))
    elif tag_name == "p":
        render_paragraph(doc, node)
    elif tag_name == "div":
        formula = node.get_text(strip=True)
        if formula:
            insert_math_block(doc, formula)
    elif tag_name == "pre":
        render_code_block(doc, node)
    elif tag_name == "table":
        render_table(doc, node)
    elif tag_name in ("ul", "ol"):
        render_list(doc, node)
    elif tag_name == "blockquote":
        render_blockquote(doc, node)
    elif tag_name == "hr":
        render_hr(doc)
    else:
        text = node.get_text(" ", strip=True)
        if text:
            para = doc.add_paragraph(text)
            set_para_space_after(para, PARA_SPACING_PT)
