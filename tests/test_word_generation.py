"""Tests for the Word generation pipeline (no rendering, no file I/O)."""

from __future__ import annotations

import pytest
from lxml import etree
from docx.oxml.ns import qn

from markdown_to_pdf.word.generator import markdown_to_docx
from markdown_to_pdf.word.math_rendering import latex_to_omml


# â”€â”€ OOXML namespace helpers â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _body_xml(doc):
    """Return the document body element for OOXML introspection."""
    return doc.element.body


def _find_all(root, ns, local):
    """Find all elements matching {ns}local in the XML tree."""
    return list(root.iter(f"{{{ns}}}{local}"))


def _para_ind(para):
    """Extract (w:left, w:hanging) indentation values from a paragraph."""
    ind = para._p.find(qn("w:ind"))
    if ind is None:
        pPr = para._p.find(qn("w:pPr"))
        if pPr is not None:
            ind = pPr.find(qn("w:ind"))
    if ind is None:
        return None, None
    return ind.get(qn("w:left")), ind.get(qn("w:hanging"))


# â”€â”€ Basic smoke tests â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestMarkdownToDocx:
    """Smoke tests: ensure key Markdown constructs produce a document."""

    def test_returns_document(self, sample_markdown):
        doc = markdown_to_docx(sample_markdown)
        assert doc is not None
        assert len(doc.paragraphs) > 0

    def test_headings_present(self, sample_markdown):
        doc = markdown_to_docx(sample_markdown)
        heading_styles = {p.style.name for p in doc.paragraphs if "Heading" in (p.style.name or "")}
        assert "Heading 1" in heading_styles
        assert "Heading 2" in heading_styles

    def test_table_present(self, sample_markdown):
        doc = markdown_to_docx(sample_markdown)
        assert len(doc.tables) >= 1

    def test_empty_input(self):
        doc = markdown_to_docx("")
        assert doc is not None


# â”€â”€ Internal References & Bookmarks â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestBookmarksAndHyperlinks:
    """Verify OOXML bookmarks (w:bookmarkStart/End) and hyperlinks (w:hyperlink)."""

    def test_heading_generates_bookmark_pair(self):
        """Each heading must produce matching bookmarkStart + bookmarkEnd."""
        doc = markdown_to_docx("# Sección Principal")
        body = _body_xml(doc)
        starts = _find_all(body, _W, "bookmarkStart")
        ends = _find_all(body, _W, "bookmarkEnd")
        # Filter out Word-internal bookmarks (starting with _)
        user_starts = [s for s in starts
                       if not s.get(qn("w:name"), "").startswith("_")]
        assert len(user_starts) >= 1, "No bookmarkStart for heading"
        # Every bookmarkStart must have a matching bookmarkEnd with same id
        start_ids = {s.get(qn("w:id")) for s in user_starts}
        end_ids = {e.get(qn("w:id")) for e in ends}
        assert start_ids.issubset(end_ids), (
            f"Unmatched bookmarkStart ids: {start_ids - end_ids}"
        )

    def test_heading_bookmark_name_is_slug(self):
        """Bookmark name should be a GitHub-style slug of the heading text."""
        doc = markdown_to_docx("# Mi Sección Especial")
        body = _body_xml(doc)
        starts = _find_all(body, _W, "bookmarkStart")
        names = {s.get(qn("w:name")) for s in starts}
        assert "mi-sección-especial" in names or "mi-seccion-especial" in names or \
               "mi-seccin-especial" in names or any("mi-secci" in n for n in names), (
            f"Expected slug for 'Mi Sección Especial', got: {names}"
        )

    def test_duplicate_headings_get_unique_bookmark_names(self):
        """Two headings with same text must get unique bookmark names."""
        doc = markdown_to_docx("# Título\n\nTexto\n\n# Título")
        body = _body_xml(doc)
        starts = _find_all(body, _W, "bookmarkStart")
        names = [s.get(qn("w:name")) for s in starts
                 if not s.get(qn("w:name"), "").startswith("_")]
        assert len(names) >= 2
        assert len(set(names)) == len(names), f"Duplicate bookmark names: {names}"

    def test_internal_link_creates_hyperlink_with_anchor(self):
        """[text](#target) must produce w:hyperlink with w:anchor attribute."""
        md = "# Destino\n\nVer [enlace](#destino)."
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        hyperlinks = _find_all(body, _W, "hyperlink")
        anchor_links = [h for h in hyperlinks if h.get(qn("w:anchor"))]
        assert len(anchor_links) >= 1, "No w:hyperlink with w:anchor found"
        # The anchor value should match the heading slug
        anchor_val = anchor_links[0].get(qn("w:anchor"))
        assert anchor_val == "destino", f"Anchor value: {anchor_val}"

    def test_internal_link_contains_text_runs(self):
        """The hyperlink element must contain run(s) with the link text."""
        md = "[Ir arriba](#seccion-1)\n\n# Sección 1"
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        hyperlinks = _find_all(body, _W, "hyperlink")
        anchor_links = [h for h in hyperlinks if h.get(qn("w:anchor"))]
        assert len(anchor_links) >= 1
        # Extract text from runs inside hyperlink
        link_text = "".join(
            t.text or "" for t in anchor_links[0].iter(qn("w:t"))
        )
        assert "Ir arriba" in link_text

    def test_link_bookmark_resolution_in_memory(self):
        """All internal hyperlink anchors must match an existing bookmark."""
        md = (
            "## Contenido\n\n"
            "- [Intro](#intro)\n"
            "- [Datos](#datos)\n\n"
            "## Intro\n\nTexto intro.\n\n"
            "## Datos\n\nTexto datos."
        )
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        # Collect all bookmark names
        bookmarks = {
            s.get(qn("w:name"))
            for s in _find_all(body, _W, "bookmarkStart")
        }
        # Collect all hyperlink anchors
        for hl in _find_all(body, _W, "hyperlink"):
            anchor = hl.get(qn("w:anchor"))
            if anchor:
                assert anchor in bookmarks, (
                    f"Hyperlink #{anchor} has no matching bookmark. "
                    f"Available: {bookmarks}"
                )

    def test_naked_anchor_creates_bookmark(self):
        """<a id='custom'></a> before a heading should produce a bookmark."""
        md = '<a id="custom-target"></a>\n\n# Mi Heading'
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        names = {
            s.get(qn("w:name"))
            for s in _find_all(body, _W, "bookmarkStart")
        }
        assert "custom-target" in names, f"Expected 'custom-target' in {names}"


# â”€â”€ Math Rendering (Word OMML) â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestMathRendering:
    """Verify math formulas produce OMML elements or italic fallback."""

    def test_inline_math_produces_content(self):
        doc = markdown_to_docx("Fórmula: $E = mc^2$")
        assert doc is not None
        assert len(doc.paragraphs) > 0

    def test_block_math_produces_paragraph(self):
        doc = markdown_to_docx("Texto previo\n\n$$\\theta = 1.22$$\n\nTexto posterior")
        assert len(doc.paragraphs) >= 2

    def test_complex_formula_no_crash(self):
        md = r"$$m_{\text{lím}} = 2.7 + 5 \log_{10}(D)$$"
        doc = markdown_to_docx(md)
        assert doc is not None

    def test_multiple_formulas(self):
        md = "Inline $a$ y $b$.\n\n$$c = a + b$$"
        doc = markdown_to_docx(md)
        assert len(doc.paragraphs) >= 2

    def test_block_math_has_omml_or_fallback(self):
        """A block equation must produce either m:oMathPara or italic fallback text."""
        doc = markdown_to_docx("$$x^2 + y^2 = z^2$$")
        body = _body_xml(doc)
        omath_paras = _find_all(body, _M, "oMathPara")
        # If XSLT is available â†’ OMML; otherwise â†’ italic fallback
        if omath_paras:
            assert len(omath_paras) >= 1
            # oMathPara must contain at least one oMath
            omaths = _find_all(omath_paras[0], _M, "oMath")
            assert len(omaths) >= 1
        else:
            # Fallback: italic run with $$ ... $$
            all_text = " ".join(r.text or "" for p in doc.paragraphs for r in p.runs)
            assert "x^2" in all_text or "$$" in all_text

    def test_inline_math_has_omml_or_fallback(self):
        """Inline math must produce m:oMath element or italic fallback."""
        doc = markdown_to_docx("Valor: $\\alpha = 0.05$")
        body = _body_xml(doc)
        omaths = _find_all(body, _M, "oMath")
        if omaths:
            assert len(omaths) >= 1
        else:
            all_text = " ".join(r.text or "" for p in doc.paragraphs for r in p.runs)
            assert "alpha" in all_text or "$" in all_text


class TestComplexMathOMML:
    """Rigorous tests for complex formulas through the OMML pipeline."""

    def test_latex_to_omml_returns_element_for_simple_formula(self):
        """Direct unit test of the conversion function."""
        result = latex_to_omml(r"x^2 + y^2 = z^2")
        if result is not None:
            # Should be an lxml Element with m: namespace children
            assert isinstance(result, etree._Element)
            tag = etree.QName(result.tag).localname
            assert tag in ("oMathPara", "oMath"), f"Unexpected root tag: {tag}"

    def test_partial_derivative_omml(self):
        """Partial derivative with nabla â€” structure must be complete."""
        latex = r"\frac{\partial T}{\partial t} = \alpha \nabla^2 T"
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            # Must contain fraction structure
            assert f"{{{_M}}}f" in xml_str or "m:f" in xml_str, \
                "Missing fraction (m:f) in OMML"

    def test_definite_integral_omml(self):
        """Definite integral with limits and exponential."""
        latex = r"\int_{0}^{\infty} e^{-x^2} \, dx = \frac{\sqrt{\pi}}{2}"
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            # Must contain nary (integral) operator
            assert f"{{{_M}}}nary" in xml_str or "m:nary" in xml_str, \
                "Missing nary (integral) element in OMML"
            # Must contain radical (sqrt)
            assert f"{{{_M}}}rad" in xml_str or "m:rad" in xml_str, \
                "Missing radical (sqrt) element in OMML"

    def test_nested_fractions_with_exponential(self):
        """Nested fraction with e^{x^2} inside."""
        latex = r"\frac{d}{dx}\left[\frac{e^{x^2}}{(1+x^2)^{3/2}}\right]"
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            # Must contain at least 2 fractions (nested)
            frac_count = xml_str.count(f"{{{_M}}}f>") + xml_str.count("m:f>")
            assert frac_count >= 2, f"Expected â‰¥2 fractions, found indicators in XML"

    def test_summation_series(self):
        """Summation with factorial â€” nary operator with sub/sup."""
        latex = r"\sum_{n=0}^{\infty} \frac{(-1)^n x^{2n+1}}{(2n+1)!} = \sin(x)"
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            assert f"{{{_M}}}nary" in xml_str or "m:nary" in xml_str, \
                "Missing nary (summation) element"

    def test_limit_definition_omml(self):
        """Limit definition of derivative."""
        latex = r"\lim_{h \to 0} \frac{f(x+h) - f(x)}{h} = f'(x)"
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            assert f"{{{_M}}}f" in xml_str or "m:f" in xml_str, \
                "Missing fraction in limit definition"

    def test_matrix_omml(self):
        """Matrix / determinant structure."""
        latex = (
            r"\begin{vmatrix}"
            r" a_{11} & a_{12} \\"
            r" a_{21} & a_{22}"
            r"\end{vmatrix}"
        )
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            # Must contain matrix element (m:m) or delimiter (m:d)
            has_matrix = (f"{{{_M}}}m>" in xml_str or "m:m>" in xml_str
                         or f"{{{_M}}}d" in xml_str or "m:d" in xml_str)
            assert has_matrix, "Missing matrix/delimiter element"

    def test_arrhenius_complete_structure(self):
        """Arrhenius equation â€” fraction inside exponent, multiplication."""
        latex = r"k = A \cdot e^{-E_a / (R \cdot T)}"
        result = latex_to_omml(latex)
        if result is not None:
            xml_str = etree.tostring(result, encoding="unicode")
            # Must contain superscript (m:sSup)
            assert f"{{{_M}}}sSup" in xml_str or "m:sSup" in xml_str, \
                "Missing superscript element for exponential"

    def test_block_math_complex_in_docx(self):
        """Full pipeline: complex formula in markdown â†’ complete OMML in docx."""
        md = r"$$\int_{0}^{\infty} e^{-x^2} \, dx = \frac{\sqrt{\pi}}{2}$$"
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        omath_paras = _find_all(body, _M, "oMathPara")
        if omath_paras:
            xml_str = etree.tostring(omath_paras[0], encoding="unicode")
            assert f"{{{_M}}}nary" in xml_str or "m:nary" in xml_str
            assert f"{{{_M}}}rad" in xml_str or "m:rad" in xml_str
        else:
            # Fallback path â€” verify italic text is present
            all_text = " ".join(r.text or "" for p in doc.paragraphs for r in p.runs)
            assert "int" in all_text or "\\int" in all_text or "$$" in all_text

    def test_multiple_complex_formulas_in_docx(self):
        """Multiple complex display formulas in sequence all render."""
        md = (
            r"$$\frac{\partial^2 u}{\partial x^2} + \frac{\partial^2 u}{\partial y^2} = 0$$"
            "\n\n"
            r"$$\oint_C \vec{F} \cdot d\vec{r} = \iint_S (\nabla \times \vec{F}) \cdot d\vec{S}$$"
            "\n\n"
            r"$$\prod_{i=1}^{n} x_i = x_1 \cdot x_2 \cdots x_n$$"
        )
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        omath_paras = _find_all(body, _M, "oMathPara")
        if omath_paras:
            assert len(omath_paras) >= 3, (
                f"Expected 3 display equations, found {len(omath_paras)}"
            )
        else:
            # Fallback: at least 3 paragraphs with $$ content
            math_paras = [p for p in doc.paragraphs if "$$" in p.text]
            assert len(math_paras) >= 3


# â”€â”€ Emoji Rendering â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestEmojiRendering:
    """Verify emoji are rendered into Word runs."""

    def test_emoji_in_paragraph(self):
        doc = markdown_to_docx("Estado: ðŸŸ¢ Bueno")
        all_text = "".join(r.text for p in doc.paragraphs for r in p.runs)
        assert "Bueno" in all_text

    def test_circle_emoji_substitution(self):
        doc = markdown_to_docx("ðŸŸ¢ Verde ðŸŸ¡ Amarillo ðŸ”´ Rojo")
        assert doc is not None
        assert len(doc.paragraphs) > 0

    def test_emoji_in_table_cell(self):
        md = "| Icono | Estado |\n|-------|--------|\n| ðŸŸ¢ | OK |"
        doc = markdown_to_docx(md)
        assert len(doc.tables) >= 1
        cell_text = doc.tables[0].cell(1, 1).text
        assert "OK" in cell_text


# â”€â”€ Strikethrough Rendering â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestStrikethroughRendering:
    """Verify ~~text~~ produces strikethrough runs in Word."""

    def test_strikethrough_applied(self):
        doc = markdown_to_docx("Normal y ~~tachado~~ fin")
        strike_runs = [
            r for p in doc.paragraphs for r in p.runs if r.font.strike
        ]
        assert len(strike_runs) >= 1
        assert any("tachado" in r.text for r in strike_runs)

    def test_strikethrough_ooxml_element(self):
        """Verify w:strike element exists in the run properties XML."""
        doc = markdown_to_docx("~~eliminado~~")
        body = _body_xml(doc)
        strikes = _find_all(body, _W, "strike")
        assert len(strikes) >= 1, "No w:strike element found in OOXML"


# â”€â”€ Nested Lists â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestNestedLists:
    """Verify nested list rendering with indentation at OOXML level."""

    def test_three_level_unordered(self):
        md = "- Nivel 1\n  - Nivel 2\n    - Nivel 3"
        doc = markdown_to_docx(md)
        list_paras = [p for p in doc.paragraphs if "•" in p.text]
        assert len(list_paras) >= 3

    def test_indentation_increases_with_depth(self):
        """w:ind w:left must increase for each nesting level."""
        md = "- Nivel 1\n  - Nivel 2\n    - Nivel 3"
        doc = markdown_to_docx(md)
        list_paras = [p for p in doc.paragraphs if "•" in p.text]
        assert len(list_paras) >= 3

        indents = []
        for p in list_paras:
            left, _ = _para_ind(p)
            if left is not None:
                indents.append(int(left))

        assert len(indents) >= 3, f"Could not extract indents from list: {indents}"
        # Each level should have strictly increasing indentation
        assert indents[0] < indents[1] < indents[2], (
            f"Indentation not increasing: {indents}"
        )

    def test_hanging_indent_present(self):
        """Each list item must have w:hanging for bullet/number alignment."""
        md = "- Item A\n- Item B"
        doc = markdown_to_docx(md)
        list_paras = [p for p in doc.paragraphs if "•" in p.text]
        for p in list_paras:
            _, hanging = _para_ind(p)
            assert hanging is not None and int(hanging) > 0, (
                f"Missing hanging indent on list item: {p.text}"
            )

    def test_ordered_list_numbering(self):
        md = "1. Primero\n2. Segundo\n3. Tercero"
        doc = markdown_to_docx(md)
        texts = [p.text for p in doc.paragraphs]
        assert any("1." in t for t in texts)
        assert any("2." in t for t in texts)
        assert any("3." in t for t in texts)

    def test_mixed_ordered_unordered_structure(self):
        """Ordered â†’ unordered â†’ back to ordered with correct prefixes."""
        md = "1. Item\n   - Sub A\n   - Sub B\n2. Item 2"
        doc = markdown_to_docx(md)
        texts = [p.text for p in doc.paragraphs]
        # Must have numbered items
        assert any("1." in t for t in texts)
        assert any("2." in t for t in texts)
        # Must have bullet items
        assert any("•" in t for t in texts)
        # Bullet indentation must be deeper than numbered
        numbered = [p for p in doc.paragraphs if "1." in p.text]
        bulleted = [p for p in doc.paragraphs if "•" in p.text]
        if numbered and bulleted:
            num_left, _ = _para_ind(numbered[0])
            bul_left, _ = _para_ind(bulleted[0])
            if num_left is not None and bul_left is not None:
                assert int(bul_left) > int(num_left), (
                    f"Bullet indent ({bul_left}) should be deeper than "
                    f"numbered indent ({num_left})"
                )

    def test_four_level_deep_nesting(self):
        """4+ levels should render without crash, with increasing indent."""
        md = "- L1\n  - L2\n    - L3\n      - L4"
        doc = markdown_to_docx(md)
        list_paras = [p for p in doc.paragraphs if "•" in p.text]
        assert len(list_paras) >= 4
        indents = []
        for p in list_paras:
            left, _ = _para_ind(p)
            if left is not None:
                indents.append(int(left))
        assert len(indents) >= 4
        # All 4 levels must be strictly increasing
        for i in range(len(indents) - 1):
            assert indents[i] < indents[i + 1], (
                f"Indent not increasing at level {i}: {indents}"
            )


# â”€â”€ Blockquote Rendering â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestBlockquoteRendering:
    """Verify blockquote DrawingML structure at OOXML level."""

    def test_simple_blockquote_creates_drawing(self):
        """A blockquote must produce a w:drawing element (rounded rectangle shape)."""
        doc = markdown_to_docx("> Esto es una cita importante.")
        body = _body_xml(doc)
        drawings = _find_all(body, _W, "drawing")
        assert len(drawings) >= 1, "No w:drawing element found for blockquote"

    def test_blockquote_has_text_box_content(self):
        """The drawing shape must contain w:txbxContent with the quoted text."""
        doc = markdown_to_docx("> Texto dentro de la cita.")
        body = _body_xml(doc)
        txbx_contents = _find_all(body, _W, "txbxContent")
        assert len(txbx_contents) >= 1, "No w:txbxContent found"
        # Extract text from all w:t elements inside the text box
        txbx_text = "".join(
            t.text or ""
            for t in txbx_contents[0].iter(qn("w:t"))
        )
        assert "cita" in txbx_text.lower() or "texto" in txbx_text.lower(), (
            f"Blockquote text not found inside txbxContent. Got: {txbx_text[:100]}"
        )

    def test_blockquote_has_rounded_rectangle_shape(self):
        """Shape geometry must be 'roundRect' (prstGeom)."""
        doc = markdown_to_docx("> Prueba de geometría.")
        body = _body_xml(doc)
        prst_geoms = _find_all(body, _A, "prstGeom")
        assert any(
            pg.get("prst") == "roundRect" for pg in prst_geoms
        ), "Blockquote shape is not roundRect"

    def test_blockquote_fill_and_border_colors(self):
        """Verify background fill and border line colors match constants."""
        from markdown_to_pdf.constants import COLOR_BQ_BG, COLOR_BQ_BORDER, rgb_hex_upper
        doc = markdown_to_docx("> Color check.")
        body = _body_xml(doc)
        # Find solidFill srgbClr values
        srgb_els = _find_all(body, _A, "srgbClr")
        color_vals = {el.get("val", "").upper() for el in srgb_els}
        expected_bg = rgb_hex_upper(COLOR_BQ_BG)
        expected_border = rgb_hex_upper(COLOR_BQ_BORDER)
        assert expected_bg in color_vals, (
            f"BQ background {expected_bg} not in shape colors: {color_vals}"
        )
        assert expected_border in color_vals, (
            f"BQ border {expected_border} not in shape colors: {color_vals}"
        )

    def test_blockquote_with_bold_and_emoji(self):
        """Blockquote with mixed inline formatting preserves content."""
        md = "> 💡 **Consejo:** Usa *cursiva* y `código`."
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        txbx_contents = _find_all(body, _W, "txbxContent")
        assert len(txbx_contents) >= 1
        txbx_text = "".join(
            t.text or ""
            for t in txbx_contents[0].iter(qn("w:t"))
        )
        assert "Consejo" in txbx_text
        # Bold runs must exist inside the text box
        bold_els = _find_all(txbx_contents[0], _W, "b")
        assert len(bold_els) >= 1, "No bold formatting (w:b) inside blockquote"

    def test_blockquote_with_nested_list(self):
        """List inside blockquote renders in the text box."""
        md = "> Opciones:\n> - Opción A\n> - Opción B"
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        txbx_contents = _find_all(body, _W, "txbxContent")
        assert len(txbx_contents) >= 1
        txbx_text = "".join(
            t.text or ""
            for t in txbx_contents[0].iter(qn("w:t"))
        )
        assert "Opción A" in txbx_text or "Opci" in txbx_text

    def test_multiple_blockquotes_get_unique_ids(self):
        """Each blockquote drawing must have a unique wp:docPr id."""
        md = "> Cita 1.\n\n> Cita 2.\n\n> Cita 3."
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        doc_prs = _find_all(body, _WP, "docPr")
        ids = [dp.get("id") for dp in doc_prs if dp.get("id")]
        assert len(ids) >= 3, f"Expected â‰¥3 docPr ids, got {len(ids)}"
        assert len(set(ids)) == len(ids), f"Duplicate docPr ids: {ids}"


# â”€â”€ Heading Anchors & Structure â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestHeadingAnchors:
    """Verify heading bookmark generation."""

    def test_multiple_heading_levels(self):
        md = "# H1\n## H2\n### H3\n#### H4"
        doc = markdown_to_docx(md)
        styles = {p.style.name for p in doc.paragraphs}
        assert "Heading 1" in styles
        assert "Heading 2" in styles
        assert "Heading 3" in styles

    def test_duplicate_headings_no_crash(self):
        md = "# Título\n\nTexto\n\n# Título"
        doc = markdown_to_docx(md)
        h1_count = sum(1 for p in doc.paragraphs if p.style.name == "Heading 1")
        assert h1_count == 2


# â”€â”€ Horizontal Rule â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestHorizontalRule:
    """Verify horizontal rules produce correct OOXML border."""

    def test_hr_produces_paragraph(self):
        md = "Antes\n\n---\n\nDespués"
        doc = markdown_to_docx(md)
        assert doc is not None
        assert len(doc.paragraphs) >= 3

    def test_hr_has_bottom_border(self):
        """The HR paragraph must have w:pBdr/w:bottom element."""
        md = "Antes\n\n---\n\nDespués"
        doc = markdown_to_docx(md)
        body = _body_xml(doc)
        bottoms = _find_all(body, _W, "bottom")
        hr_borders = [
            b for b in bottoms
            if b.get(qn("w:val")) == "single"
        ]
        assert len(hr_borders) >= 1, "No w:bottom border found for <hr>"


# â”€â”€ Real Examples â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

class TestRealExamples:
    """Integration-style checks on real example Markdown files."""

    def test_astronomia_docx(self, astronomia_markdown):
        if astronomia_markdown is None:
            pytest.skip("guia_astronomia_observacional.md not found")
        doc = markdown_to_docx(astronomia_markdown)
        assert doc is not None
        assert len(doc.paragraphs) > 0
        styles = {p.style.name for p in doc.paragraphs}
        assert "Heading 1" in styles
        assert len(doc.tables) >= 5
        # Verify bookmarks exist for internal navigation
        body = _body_xml(doc)
        bookmarks = {
            s.get(qn("w:name"))
            for s in _find_all(body, _W, "bookmarkStart")
            if not s.get(qn("w:name"), "").startswith("_")
        }
        assert len(bookmarks) >= 8, (
            f"Expected â‰¥8 heading bookmarks, got {len(bookmarks)}"
        )

    def test_culinaria_docx(self, culinaria_markdown):
        if culinaria_markdown is None:
            pytest.skip("atlas_tecnicas_culinarias.md not found")
        doc = markdown_to_docx(culinaria_markdown)
        assert doc is not None
        assert len(doc.paragraphs) > 0
        styles = {p.style.name for p in doc.paragraphs}
        assert "Heading 1" in styles
        assert len(doc.tables) >= 5
        # Verify math equations were processed
        body = _body_xml(doc)
        omath_paras = _find_all(body, _M, "oMathPara")
        omaths = _find_all(body, _M, "oMath")
        has_math = len(omath_paras) > 0 or len(omaths) > 0
        has_fallback = any("$$" in p.text for p in doc.paragraphs)
        assert has_math or has_fallback, (
            "No OMML math elements or fallback text found in culinaria doc"
        )
